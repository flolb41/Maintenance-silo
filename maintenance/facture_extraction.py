import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile

try:
    from PIL import Image, ImageFilter, ImageOps
except Exception:  # pragma: no cover - dépendance optionnelle
    Image = None
    ImageFilter = None
    ImageOps = None

try:
    import pytesseract
except Exception:  # pragma: no cover - dépendance optionnelle
    pytesseract = None

try:
    from pypdf import PdfReader
except Exception:  # pragma: no cover - dépendance optionnelle
    PdfReader = None

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover - dépendance optionnelle
    DocxDocument = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover - dépendance optionnelle
    load_workbook = None

try:
    from pdf2image import convert_from_bytes
except Exception:  # pragma: no cover - dépendance optionnelle
    convert_from_bytes = None


DATE_PATTERNS = (
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d.%m.%Y",
    "%d/%m/%y",
    "%Y-%m-%d",
    "%Y/%m/%d",
    "%d %B %Y",
    "%d %b %Y",
    "%d %m %Y",
    "%d %m %y",
)


MONTHS_FR = (
    "janvier", "fevrier", "fevrier", "mars", "avril", "mai", "juin",
    "juillet", "aout", "septembre", "octobre", "novembre", "decembre",
)
MONTHS_FR_PATTERN = (
    r"(?:janvier|février|mars|avril|mai|juin|juillet|août|septembre|octobre|novembre|décembre)"
)
MONTHS_FR_MAP = {
    "janvier": 1, "fevrier": 2, "février": 2, "mars": 3, "avril": 4,
    "mai": 5, "juin": 6, "juillet": 7, "aout": 8, "août": 8,
    "septembre": 9, "octobre": 10, "novembre": 11, "decembre": 12, "décembre": 12,
}


def _normalize_whitespace(value):
    return re.sub(r"\s+", " ", value or "").strip()


def _parse_decimal(value):
    if value is None:
        return None
    text = str(value).strip()
    text = re.sub(r"[€\sEUR]", "", text, flags=re.IGNORECASE)
    text = text.replace("\xa0", "")
    text = text.replace("\u202f", "")
    text = text.replace("−", "-")
    last_comma = text.rfind(",")
    last_dot = text.rfind(".")
    if last_comma == -1 and last_dot == -1:
        try:
            return Decimal(text)
        except InvalidOperation:
            return None
    if last_comma > last_dot:
        text = text.replace(".", "").replace(",", ".")
    elif last_dot > last_comma:
        text = text.replace(",", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_date(value):
    if value is None:
        return None
    text = _normalize_whitespace(value)
    text = re.sub(r"^.*?date\s*[:\-]\s*", "", text,
                  flags=re.IGNORECASE).strip()
    for fmt in DATE_PATTERNS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    match = re.search(
        r"(\d{1,2})\s+(" + MONTHS_FR_PATTERN + r")\s+(\d{4})", text, re.IGNORECASE)
    if match:
        month = MONTHS_FR_MAP.get(match.group(2).lower())
        if month:
            try:
                return datetime.strptime(f"{match.group(1)} {month} {match.group(3)}", "%d %m %Y").date()
            except ValueError:
                pass
    return None


def _extract_number(text):
    header = "\n".join(text.splitlines()[:30])
    patterns = [
        r"(?im)\b(?:facture|avoir)\s+n(?:[°ºo0]|um[ée]ro)?\s*[:#\-]?\s*([a-z0-9][a-z0-9._\-/]{2,39})",
        r"(?im)\bn(?:[°ºo0]|um[ée]ro)?\s*(?:de\s+)?(?:facture|avoir)\s*[:#\-]?\s*([a-z0-9][a-z0-9._\-/]{2,39})",
        r"(?im)\b(?:invoice|credit\s+note)\s*(?:number|no|#)\s*[:#\-]?\s*([a-z0-9][a-z0-9._\-/]{2,39})",
        r"(?im)\br[ée]f(?:[ée]rence)?\s+facture\s*[:#\-]?\s*([a-z0-9][a-z0-9._\-/]{2,39})",
        r"(?im)^\s*(?:facture|invoice|avoir|credit\s+note)\s*[:#\-]?\s*([a-z]{1,8}[-_/\.]?\d[a-z0-9._\-/]{2,39})",
    ]
    for pattern in patterns:
        found = re.search(pattern, header)
        if found:
            value = found.group(1).strip(" .|;:")
            value = re.sub(r"\s+", "", value)
            if (
                value.lower() not in {"date", "client", "total"}
                and _parse_date(value) is None
            ):
                return value

    lines = [_normalize_whitespace(line)
             for line in text.splitlines() if line.strip()]
    for index, line in enumerate(lines[:30]):
        if not (
            re.search(r"(?i)\bnum[ée]ro\b", line)
            and re.search(r"(?i)\bdate\b", line)
        ):
            continue
        for values_line in lines[index + 1:index + 5]:
            date_match = re.search(
                r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}",
                values_line,
            )
            if not date_match:
                continue
            before_date = values_line[:date_match.start()].strip(" |;:-")
            candidates = re.findall(
                r"(?i)\b[a-z0-9][a-z0-9._\-/]{2,39}\b",
                before_date,
            )
            for candidate in candidates:
                if not re.fullmatch(r"(?i)FR\d{11}", candidate):
                    return candidate

    if re.search(r"(?i)\b(?:facture|avoir|invoice|credit\s+note)\b", header):
        standalone_number = re.search(
            r"(?im)^\s*n(?:[°ºo0]|um[ée]ro)?\s*[:#\-]?\s*"
            r"([a-z0-9][a-z0-9._\-/]{2,39})\s*$",
            header,
        )
        if standalone_number:
            value = standalone_number.group(1).strip(" .|;:")
            if _parse_date(value) is None:
                return value

    label_pattern = re.compile(
        r"(?i)\b(?:facture|invoice|avoir|credit\s+note|n(?:[°o]|um[ée]ro)|r[ée]f[ée]rence)\b"
    )
    token_pattern = re.compile(
        r"(?i)\b[a-z]{1,8}[-_/\.]*\d[a-z0-9._\-/]{2,39}\b")
    for index, line in enumerate(lines[:30]):
        if not label_pattern.search(line):
            continue
        window = " ".join(lines[index:index + 3])
        candidates = token_pattern.findall(window)
        if candidates:
            return candidates[-1]
    return None


def _supplier_from_email(text):
    generic_domains = {
        "gmail", "googlemail", "hotmail", "outlook", "live", "orange",
        "wanadoo", "yahoo", "icloud", "laposte", "free", "sfr",
    }
    generic_parts = {
        "admin", "administration", "accueil", "contact", "compta",
        "comptabilite", "facturation", "facture", "info", "mail",
        "service", "secretariat",
    }
    emails = re.findall(
        r"(?i)\b[a-z0-9.!#$%&'*+/=?^_`{|}~-]+@(?:[a-z0-9-]+\.)+[a-z]{2,}\b",
        "\n".join(text.splitlines()[:20]),
    )
    for email in emails:
        local_part, domain = email.lower().rsplit("@", 1)
        domain_parts = domain.split(".")
        domain_name = domain_parts[-2] if len(domain_parts) >= 2 else ""
        source = local_part if domain_name in generic_domains else domain_name
        words = [
            word for word in re.split(r"[._+\-]+", source)
            if len(word) > 1 and word not in generic_parts
        ]
        if words:
            return " ".join(word.capitalize() for word in words)[:90]
    return None


def _extract_supplier(text):
    raw_lines = [line.rstrip() for line in text.splitlines() if line.strip()]
    lines = [line.strip() for line in raw_lines]
    for line in lines[:20]:
        found = re.match(
            r"(?i)^\s*(?:fournisseur|supplier|vendeur|[ée]metteur)\s*[:\-]\s*(.{2,90})$",
            line,
        )
        if found:
            value = found.group(1).strip(" -:;.")
            if "@" in value:
                return _supplier_from_email(value) or value
            return value

    for email_index, line in enumerate(lines[:20]):
        if "@" not in line:
            continue
        for candidate in reversed(raw_lines[max(0, email_index - 6):email_index]):
            candidate = re.split(r"\s{3,}", candidate.strip(), maxsplit=1)[0]
            lowered = candidate.lower()
            if (
                not candidate
                or re.search(r"\d{4,}|@", candidate)
                or re.fullmatch(r"[\d\s+.()\-]+", candidate)
                or any(token in lowered for token in (
                    "rue", "avenue", "boulevard", "route", "cedex",
                    "tél", "tel", "facture", "client",
                ))
            ):
                continue
            if re.search(r"[a-zA-ZÀ-ÖØ-öø-ÿ]", candidate):
                return candidate.strip(" -:;.")[:90]

    legal_forms = r"SARL|SA|SAS|SASU|EURL|SCI|SNC|SCA|SELARL|SCOP|EI"
    legal_name = re.compile(
        rf"(?i)^(?:(?:{legal_forms})\s+.+|.+\s+(?:{legal_forms}))$"
    )
    buyer_label = re.compile(
        r"(?i)\b(?:client|factur[ée]\s+[àa]|livr[ée]\s+[àa]|acheteur|destinataire)\b"
    )
    issuer_evidence = re.compile(
        r"(?i)\b(?:siret|siren|rcs|rm|ape|tva\s+intracommunautaire|"
        r"capital\s+social|iban)\b"
    )
    legal_candidates = [
        (index, line)
        for index, line in enumerate(lines[:20])
        if legal_name.match(line) and len(line) <= 90 and not buyer_label.search(line)
    ]
    scored_candidates = []
    for candidate_index, (line_index, line) in enumerate(legal_candidates):
        next_index = (
            legal_candidates[candidate_index + 1][0]
            if candidate_index + 1 < len(legal_candidates)
            else min(len(lines), line_index + 6)
        )
        company_block = " ".join(lines[line_index:next_index])
        preceding_block = " ".join(lines[max(0, line_index - 2):line_index])
        score = 30 - line_index
        if issuer_evidence.search(company_block):
            score += 40
        if re.search(r"(?i)\b(?:courriel|e-?mail|t[ée]l(?:[ée]phone)?)\b|@", company_block):
            score += 8
        if buyer_label.search(preceding_block):
            score -= 45
        scored_candidates.append((score, -line_index, line))
    if scored_candidates:
        return max(scored_candidates)[2].strip(" -:;.")

    email_supplier = _supplier_from_email(text)
    if email_supplier:
        return email_supplier

    stopwords = {
        "facture", "invoice", "total", "tva", "date", "numero", "n°", "client",
        "siret", "taux", "montant", "ht", "ttc", "net", "payer", "reference",
        "réf", "ref", "adresse", "email", "tel", "telephone", "code", "iban",
        "page", "page 1", "page 2", "page 3",
    }
    for line in lines[:8]:
        if len(line) < 3 or len(line) > 90:
            continue
        lowered = line.lower()
        if any(token in lowered for token in stopwords):
            continue
        if re.search(r"[a-zA-ZÀ-ÖØ-öø-ÿ]", line) and not re.search(r"\d{4,}", line):
            return line.strip(" -:;.")
    return None


AMOUNT_PATTERN = re.compile(
    r"(?<!\d)[-−]?(?:\d{1,3}(?:[ .'’\u00a0\u202f]\d{3})+|\d+)[,.]\d{2}(?!\d)"
)


def _invoice_text_quality(text):
    normalized = _normalize_whitespace(text).lower()
    signals = sum((
        bool(re.search(r"\b(?:facture|invoice)\b", normalized)),
        bool(re.search(
            r"\b(?:total|montant|net)\s+(?:ht|ttc|à\s+payer|a\s+payer)\b", normalized)),
        bool(re.search(r"\b(?:tva|vat)\b", normalized)),
        bool(re.search(r"\b(?:date|émission|emission)\b", normalized)),
        bool(AMOUNT_PATTERN.search(normalized)),
    ))
    return len(normalized), signals


def _best_invoice_text(*candidates):
    return max(
        (text for text in candidates if text),
        key=lambda text: (
            _invoice_text_quality(text)[1],
            _invoice_text_quality(text)[0],
        ),
        default="",
    )


def _needs_complementary_ocr(text):
    length, signals = _invoice_text_quality(text)
    return length < 120 or signals < 3


def _last_amount(line):
    values = AMOUNT_PATTERN.findall(line)
    return _parse_decimal(values[-1]) if values else None


ACCOUNTING_LABEL_PATTERN = re.compile(
    r"(?i)\b(?:total|montant|sous[- ]?total|base)\s+(?:de\s+)?(?:ht|tva|ttc)\b"
    r"|\b(?:net|reste|total)\s+[àa]\s+payer\b"
    r"|\b(?:acompte|d[ée]j[àa]\s+r[ée]gl[ée]|avoir)\b"
    r"|\btva(?:\s+collect[ée]e)?\b"
)


def _amount_for_label(line, label_match):
    value_end = len(line)
    next_label = ACCOUNTING_LABEL_PATTERN.search(line, label_match.end())
    if next_label:
        value_end = next_label.start()
    separator = re.search(r"[|;]", line[label_match.end():value_end])
    if separator:
        value_end = label_match.end() + separator.start()
    return _last_amount(line[label_match.end():value_end])


def _extract_tabular_totals(lines):
    for index, line in enumerate(lines):
        tax_columns = []
        for field, pattern in (
            ("montant_ht", r"(?i)\bbase\s+ht\b"),
            ("taux_tva", r"(?i)\btaux(?:\s+(?:de\s+)?tva)?\b"),
            ("montant_tva", r"(?i)\bmont(?:ant)?\.?\s+(?:de\s+)?tva\b"),
        ):
            match = re.search(pattern, line)
            if match:
                tax_columns.append((match.start(), field))
        if len(tax_columns) != 3:
            continue
        tax_columns.sort()
        bases = []
        taxes = []
        rates = []
        for values_line in lines[index + 1:index + 7]:
            if ACCOUNTING_LABEL_PATTERN.search(values_line):
                break
            values = [
                _parse_decimal(match.group(0))
                for match in AMOUNT_PATTERN.finditer(values_line)
            ]
            values = [value for value in values if value is not None]
            if len(values) < 3:
                continue
            row = {
                field: value
                for (_, field), value in zip(tax_columns, values[:3])
            }
            rate = row["taux_tva"]
            if not Decimal(0) <= rate <= Decimal(100):
                continue
            rates.append(rate)
            bases.append(row["montant_ht"])
            taxes.append(row["montant_tva"])
        if bases and taxes:
            result = {
                "montant_ht": sum(bases, Decimal(0)),
                "montant_tva": sum(taxes, Decimal(0)),
            }
            distinct_rates = {rate for rate in rates if rate is not None}
            if len(distinct_rates) == 1:
                result["taux_tva"] = distinct_rates.pop()
            return result

    field_patterns = {
        "montant_ht": r"\b(?:total|montant)\s+ht\b",
        "montant_tva": r"\b(?:montant|total)\s+(?:de\s+)?tva\b|\btva\b",
        "montant_ttc": r"\b(?:total|montant)\s+ttc\b|\bnet\s+[àa]\s+payer\b",
    }
    for index, line in enumerate(lines):
        columns = []
        for field, pattern in field_patterns.items():
            match = re.search(pattern, line, re.IGNORECASE)
            if match:
                columns.append((match.start(), field))
        if len(columns) < 2:
            continue

        columns.sort()
        for values_line in lines[index + 1:index + 3]:
            values = [
                (match.start(), _parse_decimal(match.group(0)))
                for match in AMOUNT_PATTERN.finditer(values_line)
            ]
            if len(values) != len(columns) or any(value is None for _, value in values):
                continue

            return {
                field: value
                for (_, field), (_, value) in zip(columns, values)
            }
    return {}


def _extract_amounts(text):
    lines = [_normalize_whitespace(line)
             for line in text.splitlines() if line.strip()]
    totals_zone = lines[-35:]
    result = _extract_tabular_totals(totals_zone)
    derived_fields = set()

    labels = {
        "montant_ttc": (
            (r"\btotal\s+ttc\b", r"\bmontant\s+ttc\b"),
            (r"\bnet\s+[àa]\s+payer\b", r"\btotal\s+[àa]\s+payer\b"),
        ),
        "montant_ht": (
            (r"\btotal\s+ht\b", r"\bmontant\s+ht\b"),
            (r"\bsous[- ]?total\s+ht\b", r"\bbase\s+ht\b"),
        ),
        "montant_tva": (
            (r"\bmontant\s+(?:de\s+)?tva\b", r"\btotal\s+tva\b"),
            (r"^\s*(?:dont\s+)?tva\b",),
        ),
    }
    for field, pattern_groups in labels.items():
        if field in result:
            continue
        for patterns in pattern_groups:
            for index in range(len(totals_zone) - 1, -1, -1):
                line = totals_zone[index]
                label_match = next(
                    (
                        match
                        for pattern in patterns
                        if (match := re.search(pattern, line, re.IGNORECASE))
                    ),
                    None,
                )
                if label_match is None:
                    continue
                amount = _amount_for_label(line, label_match)
                if amount is None:
                    for neighbor in totals_zone[index + 1:index + 3]:
                        if ACCOUNTING_LABEL_PATTERN.search(neighbor):
                            break
                        amount = _last_amount(neighbor)
                        if amount is not None:
                            break
                if amount is not None:
                    result[field] = amount
                    break
            if field in result:
                break

    for line in reversed(totals_zone):
        if "tva" not in line.lower():
            continue
        rates = re.findall(r"(?<!\d)(\d{1,2}(?:[,.]\d{1,2})?)\s*%", line)
        if rates:
            result["taux_tva"] = _parse_decimal(rates[-1])
            break

    ht = result.get("montant_ht")
    tva = result.get("montant_tva")
    ttc = result.get("montant_ttc")
    if ht is not None and ttc is not None:
        expected_tva = ttc - ht
        tolerance = max(Decimal("0.02"), abs(ttc) * Decimal("0.001"))
        expected_rate = (
            (abs(expected_tva) * Decimal(100) / abs(ht))
            if ht else None
        )
        common_french_rates = (
            Decimal(0), Decimal("2.10"), Decimal("5.50"),
            Decimal("10"), Decimal("20"),
        )
        coherent_common_rate = expected_rate is not None and any(
            abs(expected_rate - rate) <= Decimal("0.15")
            for rate in common_french_rates
        )
        if (
            abs(expected_tva) > tolerance
            and (tva == 0 or (
                tva is not None
                and abs(tva - expected_tva) > tolerance
                and coherent_common_rate
            ))
        ):
            result["montant_tva"] = expected_tva
            tva = expected_tva
            derived_fields.add("montant_tva")
    if ht is not None and tva is not None and ttc is None:
        result["montant_ttc"] = ht + tva
        derived_fields.add("montant_ttc")
    elif ht is not None and ttc is not None and tva is None and ttc >= ht:
        result["montant_tva"] = ttc - ht
        derived_fields.add("montant_tva")
    elif tva is not None and ttc is not None and ht is None and ttc >= tva:
        result["montant_ht"] = ttc - tva
        derived_fields.add("montant_ht")

    detected_rate = result.get("taux_tva")
    if ht and tva is not None and (
        detected_rate is None
        or (detected_rate == 0 and tva != 0)
    ):
        result["taux_tva"] = (tva * Decimal("100") /
                              ht).quantize(Decimal("0.01"))
        derived_fields.add("taux_tva")
    return result, derived_fields


def _extract_date(text):
    lines = text.splitlines()
    preferred_labels = (
        r"(?i)\b(?:date\s+(?:de\s+)?facture|date\s+d['’]?[ée]mission|"
        r"date\s+d['’]?[ée]tablissement|date\s+d['’]?[ée]dition|invoice\s+date)\b"
    )
    date_value_pattern = re.compile(
        r"\d{1,2}[./-]\d{1,2}[./-]\d{2,4}"
        r"|\d{4}[./-]\d{1,2}[./-]\d{1,2}"
        r"|\d{1,2}\s+" + MONTHS_FR_PATTERN + r"\s+\d{4}",
        re.IGNORECASE,
    )
    for index, line in enumerate(lines[:30]):
        label_match = re.search(preferred_labels, line)
        if label_match:
            candidate_text = " ".join(
                [line[label_match.end():], *lines[index + 1:index + 3]]
            )
            date_match = date_value_pattern.search(candidate_text)
            parsed = _parse_date(date_match.group(0)) if date_match else None
            if parsed:
                return parsed
    for i, line in enumerate(lines[:25]):
        candidate = line.strip()
        if i + 1 < len(lines):
            candidate = candidate + " " + lines[i + 1].strip()
        if re.search(r"(?i)\bdate\b", candidate) and not re.search(
            r"(?i)\b(?:[ée]ch[ée]ance|due|livraison)\b", candidate
        ):
            parsed = _parse_date(candidate)
            if parsed:
                return parsed

    normalized = _normalize_whitespace(text)
    for candidate in re.findall(r"(\d{1,2}[./-]\d{1,2}[./-]\d{2,4}|\d{4}[./-]\d{1,2}[./-]\d{1,2})", normalized):
        parsed = _parse_date(candidate)
        if parsed:
            return parsed
    for candidate in re.findall(r"(\d{1,2}\s+" + MONTHS_FR_PATTERN + r"\s+\d{4})", normalized, re.IGNORECASE):
        parsed = _parse_date(candidate)
        if parsed:
            return parsed
    return None


def _preprocess_image(image):
    if image is None or Image is None:
        return image
    try:
        image = image.convert("L")
        if ImageOps is not None:
            image = ImageOps.autocontrast(image, cutoff=1)
        if image.width < 1800:
            ratio = 1800 / image.width
            image = image.resize(
                (1800, int(image.height * ratio)),
                Image.Resampling.LANCZOS,
            )
        if ImageFilter is not None:
            image = image.filter(ImageFilter.SHARPEN)
        return image
    except Exception:
        return image


def _needs_header_ocr(text):
    return not (
        _extract_number(text)
        and _extract_date(text)
        and _extract_supplier(text)
    )


def _needs_totals_ocr(text):
    amounts, _ = _extract_amounts(text)
    return amounts.get("montant_ttc") is None or not (
        amounts.get("montant_ht") is not None
        or amounts.get("montant_tva") is not None
    )


def _ocr_pdf_image(image, zone=None):
    if zone and hasattr(image, "crop") and hasattr(image, "size"):
        width, height = image.size
        if zone == "header":
            image = image.crop((0, 0, width, int(height * 0.55)))
        elif zone == "totals":
            image = image.crop((0, int(height * 0.5), width, height))
    image = _preprocess_image(image)
    candidates = []
    for psm in (6, 11):
        try:
            text = pytesseract.image_to_string(
                image, config=f"--psm {psm} -l fra+eng"
            )
        except (OSError, RuntimeError):
            text = pytesseract.image_to_string(
                image, config=f"--psm {psm}"
            )
        if text:
            candidates.append(text)
        if text and not _needs_complementary_ocr(text):
            break
    return _best_invoice_text(*candidates)


def _read_pdf_text(file_obj):
    page_texts = {}
    page_count = 0
    if PdfReader is not None:
        try:
            file_obj.seek(0)
            reader = PdfReader(file_obj)
            page_count = len(reader.pages)
            selected_indices = list(range(min(page_count, 10)))
            if page_count and page_count - 1 not in selected_indices:
                selected_indices.append(page_count - 1)
            for index in selected_indices:
                page = reader.pages[index]
                try:
                    page_text = page.extract_text(
                        extraction_mode="layout") or ""
                except (TypeError, ValueError):
                    page_text = page.extract_text() or ""
                page_texts[index] = page_text
        except Exception:
            page_texts = {}
            page_count = 0

    used_ocr = False
    first_text = page_texts.get(0, "")
    last_index = max(0, page_count - 1)
    last_text = page_texts.get(last_index, "")
    target_zones = {}
    if _needs_header_ocr(first_text):
        target_zones.setdefault(0, []).append("header")
    if _needs_totals_ocr(last_text):
        target_zones.setdefault(last_index, []).append("totals")

    if target_zones and convert_from_bytes is not None and pytesseract is not None:
        try:
            file_obj.seek(0)
            content = file_obj.read()
            for index, zones in target_zones.items():
                images = convert_from_bytes(
                    content,
                    dpi=300,
                    first_page=index + 1,
                    last_page=index + 1,
                )
                ocr_text = "\n".join(
                    _ocr_pdf_image(image, zone)
                    for image in images
                    for zone in zones
                ).strip()
                if ocr_text:
                    page_texts[index] = "\n".join(
                        part for part in (
                            page_texts.get(index, "").strip(), ocr_text
                        ) if part
                    )
                    used_ocr = True
        except Exception:
            pass

    if not page_texts:
        return "", used_ocr
    return "\f".join(page_texts[index] for index in sorted(page_texts)), used_ocr


def _read_docx_text(file_obj):
    if DocxDocument is None:
        return ""
    try:
        file_obj.seek(0)
        doc = DocxDocument(file_obj)
        lines = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        lines.append(text)
        return "\n".join(lines)
    except Exception:
        return ""


def _read_xlsx_text(file_obj):
    if load_workbook is None:
        return ""
    try:
        file_obj.seek(0)
        wb = load_workbook(file_obj, read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for cell in row:
                    if cell is not None:
                        text = str(cell).strip()
                        if text:
                            lines.append(text)
        return "\n".join(lines)
    except Exception:
        return ""


def _read_image_text(file_obj):
    if Image is None or pytesseract is None:
        return ""
    try:
        file_obj.seek(0)
        image = Image.open(file_obj)
        image = _preprocess_image(image)

        def read_data(psm):
            try:
                return pytesseract.image_to_data(
                    image,
                    config=f"--psm {psm} -l fra+eng",
                    output_type=pytesseract.Output.DICT,
                )
            except (OSError, RuntimeError):
                return pytesseract.image_to_data(
                    image,
                    config=f"--psm {psm}",
                    output_type=pytesseract.Output.DICT,
                )

        def data_to_text(data):
            rows = {}
            for index, word in enumerate(data.get("text", [])):
                word = _normalize_whitespace(word)
                try:
                    confidence = float(data["conf"][index])
                except (KeyError, TypeError, ValueError):
                    confidence = -1
                if not word or confidence < 20:
                    continue
                row_key = (
                    data["block_num"][index],
                    data["par_num"][index],
                    data["line_num"][index],
                )
                try:
                    top = data["top"][index]
                except (KeyError, IndexError, TypeError):
                    top = data["line_num"][index] * 100
                rows.setdefault(row_key, []).append(
                    (data["left"][index], top, word))
            return "\n".join(
                " ".join(word for _, _, word in sorted(words))
                for words in sorted(
                    rows.values(),
                    key=lambda row: (
                        min(top for _, top, _ in row),
                        min(left for left, _, _ in row),
                    ),
                )
            )

        primary_text = data_to_text(read_data(6))
        if not _needs_complementary_ocr(primary_text):
            return primary_text
        sparse_text = data_to_text(read_data(11))
        return _best_invoice_text(primary_text, sparse_text)
    except Exception:
        return ""


def _build_detection_metadata(
    data,
    derived_fields,
    method="analyse_spatiale",
    source_text="",
):
    essential_fields = (
        "numero", "fournisseur", "date_facture", "montant_ht", "montant_ttc"
    )
    confidence = {}
    alerts = []
    for field in essential_fields:
        if data.get(field) is None:
            confidence[field] = 0
            alerts.append(f"Champ non détecté : {field}")
        elif field in derived_fields:
            confidence[field] = 75
        else:
            confidence[field] = 95

    for field in ("montant_tva", "taux_tva"):
        if data.get(field) is not None:
            confidence[field] = 75 if field in derived_fields else 90

    ht = data.get("montant_ht")
    tva = data.get("montant_tva")
    ttc = data.get("montant_ttc")
    if ht is not None and tva is not None and ttc is not None:
        tolerance = max(Decimal("0.02"), abs(ttc) * Decimal("0.005"))
        if abs((ht + tva) - ttc) > tolerance:
            alerts.append("Montants comptables incohérents")

    normalized_text = _normalize_whitespace(source_text).lower()
    signals = {
        "mot_facture": bool(re.search(
            r"\b(?:facture|invoice|avoir|credit\s+note)\b", normalized_text)),
        "libelle_comptable": bool(ACCOUNTING_LABEL_PATTERN.search(normalized_text)),
        "montants": len(AMOUNT_PATTERN.findall(normalized_text)) >= 2,
        "date_libellee": bool(re.search(
            r"\b(?:date|émission|emission|invoice\s+date)\b", normalized_text)),
    }
    score = round(sum(confidence.get(field, 0)
                  for field in essential_fields) / len(essential_fields))
    if not signals["mot_facture"]:
        alerts.append(
            "Le document ne contient pas de libellé facture clairement identifié")
        score -= 15
    if not signals["libelle_comptable"]:
        alerts.append("Aucun libellé comptable reconnu dans les totaux")
        score -= 10
    if not signals["montants"]:
        alerts.append("Nombre de montants détectés insuffisant")
        score -= 10
    if method in {"ocr_image", "pdf_texte_et_ocr"}:
        score -= 5
    if "Montants comptables incohérents" in alerts:
        score = max(0, score - 30)
    return {
        "score_global": score,
        "confiance": confidence,
        "alertes": alerts,
        "champs_calcules": sorted(derived_fields),
        "methode": method,
        "signaux": signals,
    }


def _read_office_text(file_obj, name):
    lowered = name.lower()
    if lowered.endswith(".docx"):
        return _read_docx_text(file_obj)
    if lowered.endswith(".xlsx"):
        return _read_xlsx_text(file_obj)
    return ""


def _decode_document_bytes(content):
    for encoding in ("utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return ""


def _split_document_pages(text):
    pages = [page.strip() for page in text.split("\f")]
    return [page for page in pages if page] or [text]


def extract_invoice_data_from_document(file_obj):
    if file_obj is None:
        return {}

    name = getattr(file_obj, "name", "") or ""
    lowered = name.lower()
    file_obj.seek(0)
    content = b""
    try:
        content = file_obj.read()
    except Exception:
        content = b""
    file_obj.seek(0)

    text = ""
    extraction_method = "analyse_spatiale"
    if lowered.endswith(".pdf"):
        text, used_ocr = _read_pdf_text(file_obj)
        extraction_method = "pdf_texte_et_ocr" if used_ocr else "pdf_texte"
    elif lowered.endswith((".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")):
        text = _read_image_text(file_obj)
        extraction_method = "ocr_image"
    elif lowered.endswith((".docx", ".xlsx", ".doc", ".xls")):
        text = _read_office_text(file_obj, lowered)
        if not text:
            text = _decode_document_bytes(content)

    if not text:
        text = _decode_document_bytes(content)

    if not text:
        return {}

    text = text.replace("\r", "\n")
    pages = _split_document_pages(text)
    first_page = pages[0]
    last_page = pages[-1]
    data = {}

    numero = _extract_number(first_page)
    if not numero and len(pages) > 1:
        numero = _extract_number(text)
    if numero:
        data["numero"] = numero

    fournisseur = _extract_supplier(first_page)
    if not fournisseur and len(pages) > 1:
        fournisseur = _extract_supplier(text)
    if fournisseur:
        data["fournisseur"] = fournisseur

    date_facture = _extract_date(first_page)
    if not date_facture and len(pages) > 1:
        date_facture = _extract_date(text)
    if date_facture:
        data["date_facture"] = date_facture

    amounts, derived_fields = _extract_amounts(last_page)
    data.update(amounts)
    data["_detection"] = _build_detection_metadata(
        data, derived_fields, extraction_method, source_text=text
    )

    return data
